"""Standalone worker: find files whose CONTENT (not filename) contains the
query words, across the user's real folders. This is what makes "find the
doc that mentions taxes" work -- it reads inside files.

Run as a SEPARATE PROCESS (like code_search_worker) so heavy/native
libraries (pypdf, lxml via python-docx/openpyxl) never load into the
bot's pyttsx3/COM process, and so a pathological file can be killed by the
caller's subprocess timeout instead of hanging the bot.

Usage:  python content_search_worker.py "<query>" [limit]
Prints one JSON line: a list of {path, score, snippet}, best first.
Only the LAST stdout line is trusted (extraction warnings go to stderr).
"""
import sys
import os
import re
import json
import time

HOME = os.path.expanduser("~")
ROOTS = [p for p in (os.path.join(HOME, d) for d in
         (r"OneDrive\Documents", "Downloads", "Desktop", r"OneDrive\Desktop", "Documents"))
         if os.path.isdir(p)]
SKIP = {"node_modules", ".git", "__pycache__", "site-packages", "venv", ".venv", ".cache",
        "appdata", "$recycle.bin", "obj", "bin", ".next", "dist", "build", "splice",
        "samples", "packs", ".vscode", ".idea", "library", "program files", "windows"}
STOPWORDS = {"my", "the", "a", "an", "for", "on", "in", "of", "file", "files", "find", "where",
             "is", "that", "this", "computer", "pc", "machine", "drive", "system", "laptop",
             "stuff", "thing", "things", "some", "any", "all", "me", "please", "spike", "hey",
             "check", "search", "look", "locate", "named", "called", "to", "mention", "mentions",
             "about", "document", "doc", "with", "contains", "containing", "inside"}

TEXT_EXT = {".txt", ".md", ".markdown", ".py", ".js", ".jsx", ".ts", ".tsx", ".gd", ".c", ".h",
            ".cpp", ".hpp", ".cc", ".cs", ".java", ".rb", ".php", ".go", ".rs", ".swift", ".kt",
            ".sh", ".bat", ".ps1", ".sql", ".json", ".csv", ".tsv", ".html", ".htm", ".xml",
            ".yml", ".yaml", ".ini", ".cfg", ".conf", ".toml", ".env", ".log", ".rst", ".tex",
            ".r", ".lua", ".pl", ".gradle", ".v"}
MAX_BYTES = 5 * 1024 * 1024      # skip files bigger than 5 MB
READ_CHARS = 1_000_000            # only scan the first ~1 MB of text per file
PDF_PAGE_CAP = 40                 # only read the first N pages of a PDF

# The bot's OWN generated artifacts must be excluded from search. The
# interaction log especially is a poisoned well: it records the text of
# every command, so it "contains" whatever you search for and wrongly wins
# content matches (this actually happened -- a search for "neurons"
# returned voice_interaction_log.csv because past test queries were logged
# in it). Screenshots/clips/QR are just self-generated clutter.
def _is_bot_artifact(fname):
    fl = fname.lower()
    return (fl == "voice_interaction_log.csv"
            or fl.startswith(("screenshot_", "clip_", "admin_totp_qr")))


def _keywords(q):
    words = re.findall(r"[A-Za-z0-9]+", q.lower())
    kws = [w for w in words if w not in STOPWORDS and len(w) > 1]
    return kws or [w for w in words if len(w) > 1][:3]


def _extract(path, ext):
    """Return the file's text (lowercased) or None if not extractable/failed."""
    try:
        if ext in TEXT_EXT:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(READ_CHARS).lower()
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(path)
            parts = []
            for pg in r.pages[:PDF_PAGE_CAP]:
                parts.append(pg.extract_text() or "")
            return "\n".join(parts).lower()
        if ext == ".docx":
            import docx
            d = docx.Document(path)
            return "\n".join(p.text for p in d.paragraphs).lower()
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if isinstance(cell, str):
                            parts.append(cell)
                    if len(parts) > 5000:
                        break
            wb.close()
            return " ".join(parts).lower()
    except Exception as e:
        sys.stderr.write(f"[extract fail] {path}: {e}\n")
    return None


def _snippet(text, kws, width=90):
    for kw in kws:
        i = text.find(kw)
        if i != -1:
            start = max(0, i - 30)
            return " ".join(text[start:start + width].split())
    return ""


def main():
    args = sys.argv[1:]
    query = args[0] if args else ""
    limit = int(args[1]) if len(args) > 1 else 8
    if not query:
        print(json.dumps([])); return
    kws = _keywords(query)
    if not kws:
        print(json.dumps([])); return

    extractable = TEXT_EXT | {".pdf", ".docx", ".xlsx", ".xlsm"}
    t0 = time.time()
    budget = 25.0
    hits = []
    done = False
    for root in ROOTS:
        if done:
            break
        for dirpath, dirs, files in os.walk(root):
            dirs[:] = [d for d in dirs if d.lower() not in SKIP and not d.startswith(".")]
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                if ext not in extractable or _is_bot_artifact(f):
                    continue
                full = os.path.join(dirpath, f)
                try:
                    if os.path.getsize(full) > MAX_BYTES:
                        continue
                except OSError:
                    continue
                text = _extract(full, ext)
                if not text:
                    continue
                score = sum(1 for k in kws if k in text)
                if score:
                    hits.append({"path": full, "score": score, "snippet": _snippet(text, kws)})
                if time.time() - t0 > budget:
                    done = True
                    break
            if done:
                break

    hits.sort(key=lambda h: (-h["score"], len(h["path"])))
    print(json.dumps(hits[:limit]))


if __name__ == "__main__":
    main()
